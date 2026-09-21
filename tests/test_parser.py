"""Run with:  python -m unittest discover -s tests -v"""
import datetime
import io
import unittest

from src import parser
from src.parser import (
    extract_education,
    extract_experience_months,
    extract_phone,
    extract_work_periods,
    extract_years_experience,
    format_experience,
    format_years,
    guess_name,
    parse_document,
    parse_job_description,
)

NOW = datetime.date(2026, 9, 20)

# Layout of a real fresher CV exported from LaTeX (identity details changed).
FRESHER_CV = """\
Asha Rao # asha.rao@example.com
Roll No.: 12230623055 (cid:131) +91-9876543210
Bachelor of Technology § GitHub Profile
St. Thomas’ College of Engineering and Technology (cid:239) www.linkedin.com/in/asharao
Summary
AI and ML undergraduate with strong skills in software development.
Education
Degree / Certificate University/Board CGPA / Year
Percentage
B.Tech (Artificial Intelli- Maulana Abul Kalam Azad University of 8.59 2023–2027
gence and Machine Learn- Technology
ing)
Senior Secondary (12th) CISCE Board 85.75% 2023
Secondary (10th) CISCE Board 92.8% 2021
Personal Projects
EmotionStack – Emotional Intelligence Simulator Jan-Jan 2026
A decision-based interactive web game.
Experience
• AI Wallah Virtual Internship July-July 2025
AI Wallah
– Hands-on experience in generative AI.
Technical Skills and Interests
Programming Languages:Java,Python,C
Extra-curricular Activities
– Volunteer Srey 2K25 - STCET, Kolkata April 2025
Achievements and Certification
– Software Engineer Intern Certificate -HackerRank 25th July 2025
"""

EXPERIENCED_CV = """\
Priya Menon
priya@example.com | +1 (415) 555-0134

EXPERIENCE
Senior Engineer, Acme Corp                     Jan 2020 – Present
- Built things.
Software Engineer, Globex                      06/2016 - 12/2019
Backend Intern, Initech                        Jun 2015 to Aug 2015

EDUCATION
B.Tech Computer Science, IIT Delhi             2012 – 2016
"""


class ExperienceTests(unittest.TestCase):
    def test_degree_years_are_not_experience(self):
        # The original bug: "2023–2027" under Education counted as 4 years.
        self.assertLess(extract_years_experience(FRESHER_CV, now=NOW), 0.5)

    def test_internship_with_shared_year_counts(self):
        text = "Experience\nWeb Intern, Foo  July-July 2025\nEducation\nB.Tech 2023-2027\n"
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 0.1)

    def test_only_education_means_zero(self):
        text = "Education\nB.Tech Computer Science 2019-2023\nSkills\nPython\n"
        self.assertEqual(extract_years_experience(text, now=NOW), 0.0)

    def test_jobs_are_summed_and_present_uses_today(self):
        years = extract_years_experience(EXPERIENCED_CV, now=NOW)
        # Jan 2020 -> Sep 2026 (81 mo) + Jun 2016 -> Dec 2019 (43 mo) + 3 mo internship
        self.assertAlmostEqual(years, round((81 + 43 + 3) / 12, 1))

    def test_overlapping_jobs_not_double_counted(self):
        text = "Experience\nEngineer, Foo  Jan 2020 - Dec 2021\nAnalyst, Bar  Jul 2021 - Dec 2022\n"
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 3.0)

    def test_future_dates_are_clamped(self):
        text = "Experience\nEngineer, Foo  Jan 2026 - Dec 2030\n"
        self.assertLess(extract_years_experience(text, now=NOW), 1.0)
        self.assertGreater(extract_years_experience(text, now=NOW), 0.0)

    def test_stated_claim_is_only_a_fallback(self):
        # No dated job anywhere -> the self-claim is used (degree dates still ignored).
        text = "Summary\nEngineer with 7+ years of experience\nEducation\nB.E. 2010-2014\n"
        self.assertEqual(extract_years_experience(text, now=NOW), 7.0)

    def test_dated_jobs_beat_an_inflated_claim(self):
        text = "Summary\n10 years of experience\nExperience\nEngineer, Foo Ltd  Jan 2024 - Dec 2025\n"
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 2.0)

    def test_degree_line_with_job_like_word_is_not_work(self):
        text = "John Doe\nB.E. Computer Engineer  2016 - 2020\nDeveloper, Foo Ltd  2021 - 2023\n"
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 2.0)

    def test_first_job_right_after_education_section_still_counts(self):
        text = (
            "Education\nB.Tech CSE, XYZ College  2018 - 2022\n"
            "Experience\nSoftware Engineer, Acme  Jun 2022 - May 2024\n"
        )
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 2.0)

    def test_title_on_the_line_above_the_dates(self):
        text = "Experience\nSenior Developer\nAcme Corp\nJan 2022 - Dec 2023\n"
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 2.0)

    def test_board_and_volunteer_roles_are_not_jobs(self):
        text = (
            "Experience\nManager, Foo Ltd  Jan 2020 - Dec 2020\n"
            "Leadership Roles\nMember of Board of Directors, Some Trust, 2010-Present.\n"
            "Volunteer Coordinator, Food Bank  2015 - 2019\n"
        )
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 1.0)

    def test_flattened_resume_counts_jobs_but_not_phd_research(self):
        # PDF export flattened onto one line, full-width dash, PhD research listed
        # right after the jobs (the real-world layout that used to give 14 years).
        text = (
            "Jane Roe Experience ENGINEERING INTERN 08/2016 \uff0d 12/2016 Company Name State "
            "Project: Built a model. PROCESS ENGINEER 04/2012 \uff0d 05/2013 Company Name City "
            "Improved yield with 5 equipment engineers. "
            "FLOW IN HORIZONTAL WELLS, UT Austin Aug. 2014-present. Evaluate transport. "
            "Interests Chess Education and Training May 2018 Ph.D : UT Austin Modeling Flows "
            "Additional Information VOLUNTEER Technical Editor of SPE Journal 2017-present."
        )
        # 5 + 14 months of real work only
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), round(19 / 12, 1))

    def test_work_periods_lists_only_the_jobs_that_counted(self):
        periods = extract_work_periods(EXPERIENCED_CV, now=NOW)
        self.assertEqual(sorted(p["months"] for p in periods), [3, 43, 81])

    def test_no_headings_skips_degree_lines(self):
        text = "John Doe\nB.Tech CGPA 8.1  2018 - 2022\nDeveloper, Foo Ltd  2022 - 2024\n"
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 2.0)

    def test_volunteering_and_projects_ignored(self):
        text = (
            "Experience\nDev, Foo  Jan 2024 - Jun 2024\n"
            "Projects\nApp  Jan 2020 - Dec 2023\n"
            "Extra-curricular Activities\nVolunteer  2019 - 2022\n"
        )
        self.assertAlmostEqual(extract_years_experience(text, now=NOW), 0.5)


class ContactTests(unittest.TestCase):
    def test_name_on_line_with_email(self):
        self.assertEqual(guess_name(FRESHER_CV, "fallback"), "Asha Rao")

    def test_name_never_a_heading_or_degree(self):
        self.assertEqual(guess_name("Bachelor of Technology\nSummary\nPersonal Projects", "fb"), "fb")

    def test_split_name_lines(self):
        self.assertEqual(guess_name("AHELI\nBANERJEE\nSummary", "fb"), "Aheli Banerjee")

    def test_roll_number_is_not_a_phone(self):
        self.assertEqual(extract_phone(FRESHER_CV), "+91-9876543210")

    def test_common_phone_formats(self):
        self.assertEqual(extract_phone("Call (415) 555-0134"), "(415) 555-0134")
        self.assertEqual(extract_phone("+91 98765 43210"), "+91 98765 43210")
        self.assertIsNone(extract_phone("B.Tech 2019-2023, CGPA 8.5"))


class EducationTests(unittest.TestCase):
    def test_btech_is_bachelor(self):
        self.assertEqual(extract_education(FRESHER_CV), "Bachelor's Degree")

    def test_no_substring_false_positives(self):
        for text in ("Associate Software Engineer", "NIM Master server", "mastery of Python"):
            self.assertEqual(extract_education(text), "Not detected", text)

    def test_high_school_diploma_is_high_school(self):
        self.assertEqual(extract_education("High School Diploma 2008"), "High School")

    def test_msc_and_phd_variants(self):
        self.assertEqual(extract_education("M.Sc. Physics"), "Master's Degree")
        self.assertEqual(extract_education("Ph.D : UT Austin"), "PhD / Doctorate")

    def test_job_description_uses_minimum_requirement(self):
        jd = "Bachelor's degree in CS required. Master's degree preferred."
        self.assertEqual(parse_job_description(jd)["required_education"], "Bachelor's Degree")


class JobDescriptionTests(unittest.TestCase):
    def test_singular_year_and_ranges(self):
        self.assertEqual(parse_job_description("5+ year hands-on experience")["min_years"], 5.0)
        self.assertEqual(parse_job_description("3-5 years of experience")["min_years"], 3.0)


class EndToEndTests(unittest.TestCase):
    def test_parse_txt_resume(self):
        doc = parse_document("asha.txt", FRESHER_CV.encode())
        self.assertEqual(doc.name, "Asha Rao")
        self.assertEqual(doc.phone, "+91-9876543210")
        self.assertEqual(doc.education, "Bachelor's Degree")
        self.assertLess(doc.years_experience, 0.5)

    def test_docx_tables_are_read(self):
        import docx

        d = docx.Document()
        d.add_paragraph("Experience")
        table = d.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Data Analyst, Foo"
        table.rows[0].cells[1].text = "Jan 2022 - Dec 2023"
        buf = io.BytesIO()
        d.save(buf)
        doc = parse_document("r.docx", buf.getvalue())
        self.assertIn("Jan 2022 - Dec 2023", doc.raw_text)
        self.assertAlmostEqual(doc.years_experience, 2.0)

    def test_format_years(self):
        self.assertEqual(format_years(4.0), "4")
        self.assertEqual(format_years(0.1), "0.1")
        self.assertEqual(format_years(None), "0")

    def test_format_experience_internship_and_years(self):
        # 3-month internship must show as "3 months", not "0.3 years" or "0 yrs"
        self.assertEqual(format_experience(0.25, 3), "3 months")
        self.assertEqual(format_experience(0.25, 3, compact=True), "3 mos")
        # 1 month
        self.assertEqual(format_experience(0.08, 1), "1 month")
        self.assertEqual(format_experience(0.08, 1, compact=True), "1 mo")
        # 6 months
        self.assertEqual(format_experience(0.5, 6), "6 months")
        self.assertEqual(format_experience(0.5, 6, compact=True), "6 mos")
        # 1 year
        self.assertEqual(format_experience(1.0, 12), "1 year")
        self.assertEqual(format_experience(1.0, 12, compact=True), "1 yr")
        # 1 year 3 months
        self.assertEqual(format_experience(1.25, 15), "1 year 3 months")
        self.assertEqual(format_experience(1.25, 15, compact=True), "1 yr 3 mos")
        # 2 years
        self.assertEqual(format_experience(2.0, 24), "2 years")
        self.assertEqual(format_experience(2.0, 24, compact=True), "2 yrs")
        # Zero
        self.assertEqual(format_experience(0.0, 0), "0 years")
        self.assertEqual(format_experience(0.0, 0, compact=True), "0 yrs")

    def test_internship_duration_stated_and_dated(self):
        # Stated 3-month internship
        text_stated = "Summary\nCompleted a 3 months internship at ABC Corp in React.\n"
        self.assertEqual(extract_experience_months(text_stated, now=NOW), 3)
        self.assertEqual(format_experience(extract_years_experience(text_stated, now=NOW), 3), "3 months")

        # Dated 3-month internship (June 2024 - August 2024 = 3 months)
        text_dated = "Experience\nSoftware Intern, Acme Corp  June 2024 - August 2024\n"
        months = extract_experience_months(text_dated, now=NOW)
        self.assertEqual(months, 3)
        self.assertEqual(format_experience(extract_years_experience(text_dated, now=NOW), months), "3 months")
        self.assertEqual(format_experience(extract_years_experience(text_dated, now=NOW), months, compact=True), "3 mos")

    def test_intern_word_equivalency(self):
        # '3 months intern' without the -ship suffix
        t1 = "Summary\nWorked as a 3 months intern at Startup Inc.\n"
        self.assertEqual(extract_experience_months(t1, now=NOW), 3)
        self.assertEqual(format_experience(extract_years_experience(t1, now=NOW), 3), "3 months")

        # 'intern for 6 months'
        t2 = "Experience\nWeb Developer Intern for 6 months at ABC Solutions.\n"
        self.assertEqual(extract_experience_months(t2, now=NOW), 6)
        self.assertEqual(format_experience(extract_years_experience(t2, now=NOW), 6), "6 months")

        # Heading named 'INTERN'
        t3 = "INTERN\nPython Developer, Foo Corp  Jan 2024 - Mar 2024\n"
        self.assertEqual(extract_experience_months(t3, now=NOW), 3)
        self.assertEqual(format_experience(extract_years_experience(t3, now=NOW), 3), "3 months")

        # Heading named 'SUMMER INTERN'
        t4 = "SUMMER INTERN\nResearch Intern, Bar Labs  May 2023 - July 2023\n"
        self.assertEqual(extract_experience_months(t4, now=NOW), 3)
        self.assertEqual(format_experience(extract_years_experience(t4, now=NOW), 3), "3 months")

    def test_internship_under_projects_and_other_sections(self):
        # User resume snippet under PROJECTS heading
        snippet = """PROJECTS
Full-Stack Developer Intern - KreupAl Technologies LLC May 2026 - August 2026 | Remote
• Helped develop tools to track HR compliance and company rules.
• Worked on executive dashboards for business reporting and data analysis.
• Collaborated well with the team and actively participated in daily standups.
Product & Web Developer Intern - Baha (studiobaha.com)
• Built and designed the company's Shopify website, customizing themes using Liquid templates.
• Added products to the website and managed the company's social media accounts.
• Researched competitors and the market to help the business grow.
"""
        periods = extract_work_periods(snippet, now=NOW)
        months = extract_experience_months(snippet, now=NOW)
        self.assertEqual(months, 4)
        self.assertEqual(format_experience(extract_years_experience(snippet, now=NOW), months), "4 months")
        self.assertEqual(format_experience(extract_years_experience(snippet, now=NOW), months, compact=True), "4 mos")

        # Must detect the dated internship period and the undated role
        period_names = [p["period"] for p in periods]
        self.assertIn("May 2026 - August 2026", period_names)
        self.assertTrue(any("Baha" in p["context"] for p in periods))

    def test_job_description_with_internship_months(self):
        from src.parser import parse_job_description
        jd = parse_job_description("Software Engineering Intern: Requires 3 months internship experience in Python and React.")
        self.assertEqual(jd["min_years"], 0.25)
        self.assertTrue(jd["is_internship"])

    def test_open_ended_role_reconciled_with_stated_claim(self):
        text = """
        Sourav Das
        Skills: JavaScript, 5 years of experience SQL, 5 years of experience
        Work History
        Software Developer , 12/2015 to Current Company Name
        Computer Engineer Intern , 06/2013 to 09/2013 Company Name
        """
        months = extract_experience_months(text, now=NOW)
        self.assertEqual(months, 60)
        self.assertEqual(extract_years_experience(text, now=NOW), 5.0)


if __name__ == "__main__":
    unittest.main()

