"""
Document extraction + structured parsing.

Handles: PDF / DOCX / TXT ingestion, and pulling out
name, email, phone, education, years of experience, and a skills list
from free-form resume or job-description text.

Years of experience come ONLY from paid/professional work: full-time roles,
part-time roles, contracts, freelance work and internships. A date range is
counted only when it (a) is not inside an Education / Projects / Skills /
Certifications / Volunteer / Publications section, (b) is not next to degree
or grade wording (B.Tech, M.S., Ph.D, CGPA ...), and (c) sits next to a job
title, "intern"/"trainee" wording or an employer marker. Overlapping jobs are
merged so they are never double counted. A stated "N years of experience"
phrase is used only as a last resort when the resume has no dated work roles.
"""
import bisect
import datetime as _dt
import io
import re
from dataclasses import dataclass, field

from .skills_taxonomy import MASTER_SKILLS, SYNONYMS, SKILL_TO_CATEGORY, canonicalize

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
URL_RE = re.compile(
    r"(?:https?://[^\s<>\[\]{}\"']+|www\.[^\s<>\[\]{}\"']+|(?:https?://)?(?:linkedin\.com|github\.com)/[^\s<>\[\]{}\"']+)",
    re.IGNORECASE,
)
# International-aware phone matcher:
#   - optional +country code (1-3 digits)
#   - optional parenthesized area/STD code
#   - then 2-4 groups of digits separated by space/dot/dash, 7-12 digits total
# Covers US "(555) 123-4567", Indian "+91 98765 43210" / "+91-9876543210",
# and plain unbroken 10-digit numbers, while still requiring digit boundaries
# so it doesn't grab pieces of longer numeric strings (IDs, years, etc.)
# NOTE: this is only a *candidate* finder - extract_phone() validates each
# candidate (digit count, "Roll No."-style labels) before accepting it.
PHONE_RE = re.compile(
    r"(?<![\d/])"
    r"(?:\+\d{1,3}[\s.-]?)?"
    r"(?:\(\d{2,4}\)[\s.-]?)?"
    r"\d{3,5}[\s.-]?\d{3,4}(?:[\s.-]?\d{2,4})?"
    r"(?![\d/])"
)
# Labels that mean "the number next to me is an ID, not a phone number".
PHONE_ID_LABEL_RE = re.compile(
    r"\b(?:roll|reg(?:istration)?|enrol(?:l?ment)?|student|employee|emp|id|uid|"
    r"aadhaar|aadhar|pan|passport|ssn|cgpa|gpa|pin\s?code|zip)\b[^\d]{0,12}$",
    re.IGNORECASE,
)

# "5 years of experience", "5+ yrs hands-on experience", "1.5 years experience",
# "3-5 years of professional experience" (captures the lower bound).
YEARS_EXP_RE = re.compile(
    r"(?<![\d.])(\d{1,2}(?:\.\d)?)(?:\s*(?:-|–|—|to)\s*\d{1,2})?\s*\+?\s*(?:years?|yrs?)\.?\s*(?:of\s+)?"
    r"(?:(?:relevant|professional|hands-on|industry|work|working|total|overall|proven)\s+)*"
    r"experience",
    re.IGNORECASE,
)
# Kept for backwards compatibility - experience is now computed by
# _date_ranges() below, which understands months ("Jan 2020 - Mar 2022").
DATE_RANGE_RE = re.compile(
    r"(19|20)\d{2}\s*(?:-|to|–|—)\s*(?:(19|20)\d{2}|present|current)",
    re.IGNORECASE,
)
# A bare "2019-2023"-style year range satisfies the phone digit-count shape too;
# filter those out so employment date ranges never get mistaken for a phone number.
YEAR_RANGE_LOOKALIKE_RE = re.compile(
    r"^(?:19|20)\d{2}[\s.-]*(?:(?:19|20)\d{2}|present|current)$", re.IGNORECASE
)

# ---------------------------------------------------------------------------
# Education
# ---------------------------------------------------------------------------
# Ordered highest -> lowest. Every pattern uses word boundaries so that
# "Associate Software Engineer", "Scrum Master" or "mastery" never count as a
# degree (the old substring matching did exactly that).
_EDU_PATTERNS = [
    (r"\bph\.?\s?d\b|\bdoctorate\b|\bdoctoral\b", "PhD / Doctorate"),
    (
        r"\bmaster(?:['’]?s)?\s+(?:of|in|degree)\b|\bmaster['’]s\b|\bmasters\b"
        r"|\bm\.\s?tech\b|\bmtech\b|\bm\.\s?sc\b|\bmsc\b|\bmba\b|\bmca\b"
        r"|\bm\.s\b|\bm\.e\b|\bm\.a\b",
        "Master's Degree",
    ),
    (
        r"\bbachelor|\bb\.\s?tech\b|\bbtech\b|\bb\.\s?sc\b|\bbsc\b|\bb\.s\b"
        r"|\bb\.e\b|\bb\.a\b|\bb\.\s?com\b|\bbcom\b|\bbca\b|\bbba\b",
        "Bachelor's Degree",
    ),
    (r"\bassociate(?:['’]?s)?\s+(?:degree|of)\b", "Associate Degree"),
    (r"(?<!high school )(?<!school )\bdiplomas?\b", "Diploma"),
    (
        r"\bhigh\s+school\b|\bsenior\s+secondary\b|\bhigher\s+secondary\b|\b12th\b|\bhsc\b",
        "High School",
    ),
]
EDUCATION_LEVELS = [(re.compile(p, re.IGNORECASE), label) for p, label in _EDU_PATTERNS]

# Words that mean "this date range belongs to a degree, not a job".
# Strong: a degree name or grade next to the dates -> never work experience.
_DEGREE_RE = re.compile(
    r"\b(?:cgpa|sgpa|gpa|percentage|graduation|dissertation|thesis|scholarship|fellowship)\b|"
    + "|".join(p for p, label in _EDU_PATTERNS if label != "High School"),
    re.IGNORECASE,
)
# Weak: institution / student wording -> education, unless the same line also
# carries a real job title (e.g. "Teaching Assistant, XYZ University").
_STUDENT_CONTEXT_RE = re.compile(
    r"\b(?:universit(?:y|ies)|college|institute|school|academy|campus|semester|"
    r"pursuing|undergraduate|coursework|student)\b",
    re.IGNORECASE,
)

# Unpaid roles that carry a title but are not employment.
_UNPAID_ROLE_RE = re.compile(
    r"\b(?:volunteer\w*|community\s+service|board\s+of\s+directors|board\s+member|"
    r"member\s+of\s+the\s+board)\b",
    re.IGNORECASE,
)

# Words that mean "this date range belongs to a job".
_WORK_TITLE_RE = re.compile(
    r"\b(?:intern(?:s|ship|ships)?|trainees?|apprentice(?:s|ship)?|engineers?|developers?|"
    r"dev|sde|swe|programmers?|analysts?|consultants?|managers?|executives?|associates?|"
    r"officers?|assistants?|administrators?|admin|coordinators?|specialists?|architects?|"
    r"designers?|scientists?|researchers?|technicians?|technologists?|auditors?|accountants?|"
    r"advis[eo]rs?|leads?|directors?|supervisors?|representatives?|freelanc(?:e|er|ing)|"
    r"contractors?|teachers?|lecturers?|professors?|tutors?|instructors?|trainers?|clerks?|"
    r"secretary|recruiters?|strategists?|generalists?|founders?|co-?founders?|owners?|"
    r"editors?|writers?|copywriters?|marketers?|salesman|salesperson|sales|cashiers?|"
    r"receptionists?|operators?|mechanics?|electricians?|drivers?|nurses?|pharmacists?|"
    r"chefs?|waiters?|baristas?|bankers?|testers?|president|vp|cto|ceo|coo|cfo|qa)\b",
    re.IGNORECASE,
)
_EMPLOYER_RE = re.compile(
    r"\b(?:pvt|private\s+limited|ltd|limited|inc|llc|llp|corp|corporation|technologies|"
    r"solutions|systems|labs?|software|consulting|consultancy|services|company|co|gmbh|"
    r"group|industries|enterprises|bank|studios?|startup|agency|firm|full[\s-]?time|"
    r"part[\s-]?time|contract|remote|on-?site|hybrid)\b",
    re.IGNORECASE,
)

# ---------------------------------------------------------------------------
# Resume section headings
# ---------------------------------------------------------------------------
_EXPERIENCE_HEADINGS = {
    "experience", "work experience", "professional experience", "work history",
    "employment", "employment history", "employment experience", "work",
    "professional background", "career history", "relevant experience",
    "industry experience", "industrial experience", "industrial training",
    "internship", "internships", "internship experience", "research experience",
    "teaching experience",
}
_EDUCATION_HEADINGS = {
    "education", "academic background", "academics", "academic qualifications",
    "educational qualifications", "education qualifications", "qualifications",
    "academic details", "educational background", "education and training",
}
_OTHER_HEADINGS = {
    "summary", "professional summary", "career summary", "objective",
    "career objective", "profile", "professional profile", "about me", "about",
    "projects", "personal projects", "academic projects", "key projects",
    "selected projects", "project experience", "skills", "technical skills",
    "core skills", "key skills", "interests", "areas of interest", "hobbies",
    "languages", "language", "certifications", "certification", "certificates",
    "achievements", "accomplishments", "awards", "honors", "honours",
    "extra curricular activities", "extracurricular activities",
    "co curricular activities", "activities", "volunteer experience",
    "volunteering", "volunteer work", "leadership", "leadership experience",
    "positions of responsibility", "publications", "research papers",
    "references", "declaration", "personal details", "personal information",
    "training", "courses", "coursework", "relevant coursework", "workshops",
    "highlights", "currently learning", "tools", "technologies",
    "tools and technologies", "volunteer", "additional information",
    "selected publications", "leadership roles", "community service",
}

STOPWORD_NAME_LINES = (
    "summary", "objective", "experience", "education", "skills",
    "profile", "resume", "curriculum vitae", "highlights", "accomplishments",
)

_TITLE_WORDS = {
    "engineer", "developer", "scientist", "analyst", "designer", "manager",
    "intern", "student", "consultant", "architect", "administrator",
    "specialist", "lead", "trainee",
}
# Words that never appear in a person's name but do appear in resume headers.
_NON_NAME_WORDS = _TITLE_WORDS | {
    "resume", "cv", "curriculum", "vitae", "summary", "objective", "profile",
    "contact", "email", "phone", "mobile", "tel", "address", "github",
    "linkedin", "portfolio", "roll", "no", "dob", "bachelor", "master",
    "masters", "technology", "university", "college", "institute", "school",
    "engineering", "education", "experience", "skills", "projects", "of",
    "and", "the", "for", "in", "at", "btech", "mtech", "degree",
}


@dataclass
class ParsedDocument:
    raw_text: str
    name: str = "Not detected"
    email: str = "Not detected"
    phone: str = "Not detected"
    linkedin_url: str = "Not detected"
    github_url: str = "Not detected"
    portfolio_url: str = "Not detected"
    education: str = "Not detected"
    years_experience: float = 0.0
    experience_months: int = 0
    skills: list = field(default_factory=list)
    parse_confidence: float = 1.0
    confidence_reasons: list = field(default_factory=list)
    input_hash: str = ""
    resolved_present_date: str = ""


def format_years(years, compact: bool = False) -> str:
    """Human-friendly years: 4 -> "4" (or "4 yrs"), 0.1 -> "0.1", 2.5 -> "2.5"."""
    try:
        s = f"{float(years or 0):.1f}"
    except (TypeError, ValueError):
        return "0 yrs" if compact else "0"
    val = s[:-2] if s.endswith(".0") else s
    if compact:
        suffix = "yr" if val == "1" else "yrs"
        return f"{val} {suffix}"
    return val


# ---------------------------------------------------------------------------
# File -> text
# ---------------------------------------------------------------------------
def extract_text_from_pdf(file_bytes: bytes) -> str:
    import pdfplumber
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            # x_tolerance=2 stops tightly-kerned (LaTeX) PDFs from gluing words
            # together ("Implementedpressure-basedgameplay...").
            try:
                t = page.extract_text(x_tolerance=2) or ""
            except TypeError:
                t = page.extract_text() or ""
            # Some PDFs store the destination separately from the visible link text.
            for hyperlink in getattr(page, "hyperlinks", []):
                uri = hyperlink.get("uri")
                if uri:
                    t += f"\n{uri}"
            text_chunks.append(t)
    return "\n".join(text_chunks)


def _dedupe_cells(cells):
    out = []
    for c in cells:
        c = " ".join(c.split())
        if c and (not out or out[-1] != c):  # merged cells repeat their text
            out.append(c)
    return out


def extract_text_from_docx(file_bytes: bytes) -> str:
    import docx
    from docx.table import Table
    doc = docx.Document(io.BytesIO(file_bytes))

    lines = []
    # Contact details often live in the page header.
    for section in doc.sections:
        for p in section.header.paragraphs:
            if p.text.strip():
                lines.append(p.text)

    # Walk the body in document order so resumes laid out with TABLES
    # (very common Word templates) are not silently dropped.
    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = _dedupe_cells(c.text for c in row.cells)
                if cells:
                    lines.append("  |  ".join(cells))
        else:
            lines.append(block.text)
    text = "\n".join(lines)

    # python-docx exposes hyperlink destinations through the document
    # relationships, while paragraph.text contains only their display text.
    links = []
    for relationship in doc.part.rels.values():
        if relationship.is_external and relationship.target_ref.startswith(("http://", "https://")):
            if relationship.target_ref not in links:
                links.append(relationship.target_ref)
    if links:
        text += "\n" + "\n".join(links)
    return text


def extract_text(filename: str, file_bytes: bytes) -> str:
    lower = filename.lower()
    try:
        if lower.endswith(".pdf"):
            return extract_text_from_pdf(file_bytes)
        if lower.endswith(".docx"):
            return extract_text_from_docx(file_bytes)
        # plain text fallback
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception:
        try:
            return file_bytes.decode("utf-8", errors="ignore")
        except Exception:
            return ""


# ---------------------------------------------------------------------------
# Section detection
# ---------------------------------------------------------------------------
def _classify_heading(line: str):
    """Return "experience" / "education" / "other" if `line` is a resume
    section heading, else None. Handles "Achievements and Certification",
    "Experience & Internships", trailing colons, ALL CAPS, etc."""
    if len(line) > 60:
        return None
    norm = re.sub(r"[^a-z&/ ]+", " ", line.lower().replace("-", " "))
    norm = re.sub(r"\s+", " ", norm).strip()
    if not norm or len(norm.split()) > 6:
        return None
    parts = [p.strip() for p in re.split(r"\s*&\s*|\s*/\s*|\s+and\s+", norm) if p.strip()]
    kinds = []
    for part in parts:
        if part in _EXPERIENCE_HEADINGS:
            kinds.append("experience")
        elif part in _EDUCATION_HEADINGS:
            kinds.append("education")
        elif part in _OTHER_HEADINGS:
            kinds.append("other")
        else:
            return None
    return kinds[0] if kinds else None


def _split_sections(text: str):
    """[(kind, [lines])] where kind is preamble/experience/education/other."""
    sections = []
    kind, buf = "preamble", []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        heading = _classify_heading(line)
        if heading:
            sections.append((kind, buf))
            kind, buf = heading, []
        else:
            buf.append(line)
    sections.append((kind, buf))
    return sections


def _section_text(text: str, kind: str):
    """Joined text of every section of `kind`, or None if there is none."""
    chunks = [lines for k, lines in _split_sections(text) if k == kind and lines]
    if not chunks:
        return None
    return "\n".join("\n".join(lines) for lines in chunks)


# ---------------------------------------------------------------------------
# Work-experience date ranges
# ---------------------------------------------------------------------------
_MONTH_PAT = (
    r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?"
    r"|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"
)
_MONTH_NUM = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}
_YEAR_PAT = r"(?:19|20)\d{2}"
_SEP_PAT = r"\s*(?:[-–—‒―‑−－]|\bto\b|\buntil\b|\btill\b)\s*"
_PRESENT_PAT = r"(?:present|current|now|ongoing|today|date)"

# "Jan 2020 - Mar 2022", "2019 - 2023", "2023 - Present", "Dec 2013 to Current"
_RANGE_FULL_RE = re.compile(
    rf"(?<![\w/])(?:(?P<sm>{_MONTH_PAT})\.?,?\s*)?(?P<sy>{_YEAR_PAT}){_SEP_PAT}"
    rf"(?:(?:(?P<em>{_MONTH_PAT})\.?,?\s*)?(?P<ey>{_YEAR_PAT})|(?P<pres>{_PRESENT_PAT}))(?![\w])",
    re.IGNORECASE,
)
# "July-July 2025", "Jun - Aug 2024" (the year is written once, at the end)
_RANGE_MONTHS_RE = re.compile(
    rf"(?<![\w/])(?P<sm>{_MONTH_PAT})\.?{_SEP_PAT}(?P<em>{_MONTH_PAT})\.?,?\s*(?P<y>{_YEAR_PAT})(?![\w])",
    re.IGNORECASE,
)
# "06/2013 to 09/2013", "12/2015 - Current"
_RANGE_NUMERIC_RE = re.compile(
    rf"(?<![\w/])(?P<sm>0?[1-9]|1[0-2])\s*[/.]\s*(?P<sy>{_YEAR_PAT}){_SEP_PAT}"
    rf"(?:(?P<em>0?[1-9]|1[0-2])\s*[/.]\s*(?P<ey>{_YEAR_PAT})|(?P<pres>{_PRESENT_PAT}))(?![\w])",
    re.IGNORECASE,
)


def _month_num(token):
    if not token:
        return None
    if token.isdigit():
        return int(token)
    return _MONTH_NUM[token[:3].lower()]


def _find_ranges(text: str):
    """Yield (match, (start_year, start_month, end_year, end_month, is_present))
    for every date range in `text`, without overlapping matches."""
    found = []
    for m in _RANGE_FULL_RE.finditer(text):
        found.append((m, (int(m["sy"]), _month_num(m["sm"]),
                          int(m["ey"]) if m["ey"] else None,
                          _month_num(m["em"]), bool(m["pres"]))))
    for m in _RANGE_MONTHS_RE.finditer(text):
        sm, em, y = _month_num(m["sm"]), _month_num(m["em"]), int(m["y"])
        found.append((m, (y if sm <= em else y - 1, sm, y, em, False)))
    for m in _RANGE_NUMERIC_RE.finditer(text):
        found.append((m, (int(m["sy"]), _month_num(m["sm"]),
                          int(m["ey"]) if m["ey"] else None,
                          _month_num(m["em"]), bool(m["pres"]))))

    found.sort(key=lambda item: (item[0].start(), -(item[0].end() - item[0].start())))
    last_end = -1
    for m, parts in found:
        if m.start() < last_end:
            continue
        last_end = m.end()
        yield m, parts


def _interval(sy, sm, ey, em, is_present, now):
    """Month-index interval [start, end) or None. Future dates are clamped to
    today so "2023-2027" style ranges can never inflate experience."""
    now_idx = now.year * 12 + now.month - 1
    start = sy * 12 + ((sm or 1) - 1)
    if is_present:
        end = now_idx + 1
    elif em:
        end = ey * 12 + em                 # end month is inclusive
    elif sm:
        end = ey * 12 + (sm - 1)           # "Jun 2019 - 2021": same month, 2 years on
    else:
        end = ey * 12                      # "2019 - 2023": whole years between
        if ey == sy:
            end = start + 12               # a lone-year range like "2021 - 2021"
    end = min(end, now_idx + 1)
    if end <= start or end - start > 12 * 50:
        return None
    return start, end


# Headings recognised even when a PDF/DOCX export flattened the whole resume
# onto one line ("... Experience ENGINEERING INTERN 08/2016 - 12/2016 ...").
# Only unambiguous phrases, matched in Title Case / UPPER CASE and followed by a
# capital letter or digit, so ordinary prose ("experience with Python") is safe.
_INLINE_HEADING_SKIP = {
    "work", "training", "activities", "profile", "about", "highlights", "courses",
    "tools", "technologies", "language", "languages", "leadership", "coursework",
}


def _build_inline_heading_re():
    lookup = {}
    for kind, names in (
        ("experience", _EXPERIENCE_HEADINGS),
        ("education", _EDUCATION_HEADINGS),
        ("other", _OTHER_HEADINGS),
    ):
        for name in names:
            if name in _INLINE_HEADING_SKIP:
                continue
            for variant in {name.title(), name.upper(), name.title().replace(" And ", " and ")}:
                lookup[variant] = kind
    alts = sorted((re.escape(v) for v in lookup), key=len, reverse=True)
    rx = re.compile(
        r"(?<![\w&/-])(?:" + "|".join(alts) + r")(?!\w)(?=\s*:?\s*[A-Z0-9(\u2022\u00b7\u25aa\u25cf])"
    )
    return rx, lookup


_INLINE_HEADING_RE, _INLINE_HEADING_KIND = _build_inline_heading_re()


def _section_markers(text: str):
    """Sorted [(position, kind)] for every section heading in `text`, whether
    the heading is on its own line or buried inside one long flattened line."""
    marks = {}
    for m in re.finditer(r"[^\n]+", text):
        kind = _classify_heading(m.group(0).strip())
        if kind:
            marks[m.start()] = kind
    for m in _INLINE_HEADING_RE.finditer(text):
        marks.setdefault(m.start(), _INLINE_HEADING_KIND[m.group(0)])
    return sorted(marks.items())


def _section_kind_at(markers, positions, pos: int) -> str:
    i = bisect.bisect_right(positions, pos) - 1
    return markers[i][1] if i >= 0 else "preamble"


# --- local "entry header" context around a date range ------------------------
_SENT_BREAK_RE = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\u2022\u00b7\u25aa\u25cf])")
_ABBREVIATIONS = {"inc", "ltd", "pvt", "co", "corp", "llc", "sr", "jr", "dr", "mr", "ms", "mrs", "st", "dept"}


def _real_breaks(s: str):
    """Sentence breaks, ignoring 'Inc. ' / 'Ltd. ' / initials such as 'J. Smith'."""
    for m in _SENT_BREAK_RE.finditer(s):
        prev = re.search(r"(\w+)[.!?]$", s[:m.start()])
        if prev and (prev.group(1).lower() in _ABBREVIATIONS or len(prev.group(1)) == 1):
            continue
        yield m


def _cut_before(s: str) -> str:
    last = None
    for last in _real_breaks(s):
        pass
    return s[last.end():] if last else s


def _cut_after(s: str) -> str:
    first = next(_real_breaks(s), None)
    return s[:first.start()] if first else s


def _line_bounds(text: str, start: int, end: int):
    ls = text.rfind("\n", 0, start) + 1
    le = text.find("\n", end)
    return ls, (len(text) if le == -1 else le)


def _adjacent_lines(text: str, ls: int, le: int):
    """The nearest non-blank line above and below (title/company are often on
    the line next to the dates), minus headings and lines with their own dates."""
    found = []
    pos = ls
    while pos > 0:
        pe = pos - 1
        ps = text.rfind("\n", 0, pe) + 1
        line, pos = text[ps:pe].strip(), ps
        if line:
            found.append(line[-80:])
            break
    pos = le
    while pos < len(text):
        ns = pos + 1
        ne = text.find("\n", ns)
        ne = len(text) if ne == -1 else ne
        line, pos = text[ns:ne].strip(), ne
        if line:
            found.append(line[:80])
            break
    return [
        ln for ln in found
        if not _classify_heading(ln) and next(_find_ranges(ln), None) is None
    ]


def _work_periods(text: str, now):
    """[(interval, period_text, context)] for every date range that is a real job."""
    markers = _section_markers(text)
    positions = [pos for pos, _ in markers]
    periods = []
    for m, parts in _find_ranges(text):
        # (a) never inside Education / Projects / Skills / Certifications / ...
        if _section_kind_at(markers, positions, m.start()) in ("education", "other"):
            continue

        ls, le = _line_bounds(text, m.start(), m.end())
        before = _cut_before(text[max(ls, m.start() - 120):m.start()])
        after = _cut_after(text[m.end():min(le, m.end() + 90)])
        own = f"{before} {after}"
        own_has_title = bool(_WORK_TITLE_RE.search(own))

        # (b) never next to a degree / grade / unpaid role, or an institution without a job title
        if _DEGREE_RE.search(own) or _UNPAID_ROLE_RE.search(own):
            continue
        if _STUDENT_CONTEXT_RE.search(own) and not own_has_title:
            continue

        # (c) must look like a job: title / intern / employer marker nearby
        context = own
        if le - ls <= 140:
            context += " " + " ".join(_adjacent_lines(text, ls, le))
        if not (_WORK_TITLE_RE.search(context) or _EMPLOYER_RE.search(context)):
            continue

        iv = _interval(*parts, now)
        if iv:
            snippet = re.sub(r"\s+", " ", f"{before[-40:]}{m.group(0)}{after[:30]}").strip()
            periods.append((iv, m.group(0).strip(), snippet))
    return periods


def _merged_months(intervals) -> int:
    merged = []
    for s, e in sorted(intervals):  # union, so overlapping jobs aren't double counted
        if merged and s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return sum(e - s for s, e in merged)


# A resume with no dated job at all may still say "5 years of experience".
# Set to False to ignore such self-claims completely (dated jobs only).
USE_STATED_YEARS_FALLBACK = True


def extract_work_periods(text: str, now=None, *args, **kwargs):
    """The jobs / internships that count towards experience, for verification:
    [{"period": "Jan 2020 - Present", "months": 81, "context": "..."}]"""
    now = now or _dt.date.today()
    return [
        {"period": period, "months": iv[1] - iv[0], "context": ctx}
        for iv, period, ctx in _work_periods(text, now)
    ]


def extract_years_experience(text: str, now=None) -> float:
    """Years of *work* experience (full-time, part-time, contract, internships).

    Education, projects, certifications, volunteering, publications and
    research/degree dates never count. Overlapping jobs are merged.
    Only if no dated work role is found is a stated "N years of experience"
    phrase used as a fallback (see USE_STATED_YEARS_FALLBACK).
    """
    now = now or _dt.date.today()
    months = _merged_months([iv for iv, _, _ in _work_periods(text, now)])
    if months:
        return round(months / 12, 1)
    if USE_STATED_YEARS_FALLBACK:
        claims = [float(m.group(1)) for m in YEARS_EXP_RE.finditer(text)]
        claims = [c for c in claims if 0 < c <= 50]
        if claims:
            return max(claims)
    return 0.0


# ---------------------------------------------------------------------------
# Name / phone / education
# ---------------------------------------------------------------------------
_NAME_SPLIT_RE = re.compile(r"[|•·▪◦●■◆#§¶†‡*,;]|\s{3,}")


def _header_segments(line: str):
    """Split a header line like "Sourav Das # me@x.com" into candidate chunks
    with the contact details / icon glyphs removed."""
    s = re.sub(r"\(cid:\d+\)", " | ", line)
    s = EMAIL_RE.sub(" | ", s)
    s = URL_RE.sub(" | ", s)
    s = PHONE_RE.sub(" | ", s)
    s = re.sub(r"[\ue000-\uf8ff]", " | ", s)
    return [seg.strip(" .:-–—_") for seg in _NAME_SPLIT_RE.split(s) if seg.strip(" .:-–—_")]


def _name_words_ok(words) -> bool:
    return not any(w.lower().strip(".,-:;()") in _NON_NAME_WORDS for w in words)


def _is_name_word(w: str) -> bool:
    return re.sub(r"[.'’\-]", "", w).isalpha()


def guess_name(text: str, fallback: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    # Look through the first 20 meaningful lines.
    for i, line in enumerate(lines[:20]):
        segments = _header_segments(line)
        for seg in segments:
            if len(seg) > 45 or _classify_heading(seg) or seg.lower() in STOPWORD_NAME_LINES:
                continue
            words = seg.split()
            if not _name_words_ok(words):
                continue

            # Normal case: name is on one line (possibly followed by contact info).
            if 1 < len(words) <= 4 and all(_is_name_word(w) for w in words):
                return seg.title() if seg.isupper() else seg

            # PDF-layout case: name split across two lines,
            # e.g. "AHELI" followed by "BANERJEE".
            if len(words) == 1 and seg.isalpha() and len(segments) == 1 and i + 1 < len(lines):
                nxt = _header_segments(lines[i + 1])
                if len(nxt) == 1:
                    nxt_seg = nxt[0]
                    if (
                        nxt_seg.isalpha()
                        and len(nxt_seg) <= 30
                        and _name_words_ok([nxt_seg])
                        and nxt_seg.lower() not in STOPWORD_NAME_LINES
                    ):
                        return f"{seg.title()} {nxt_seg.title()}"

    return fallback


def extract_phone(text: str):
    """First candidate that is really a phone number: right digit count, not a
    "2019-2023" date range, and not labelled as an ID ("Roll No.: 12230623055")."""
    for m in PHONE_RE.finditer(text):
        candidate = m.group(0).strip()
        digits = re.sub(r"\D", "", candidate)
        has_plus = candidate.startswith("+")

        if YEAR_RANGE_LOOKALIKE_RE.match(candidate):
            continue
        if has_plus:
            if not 8 <= len(digits) <= 15:
                continue
        else:
            plausible = (
                7 <= len(digits) <= 10
                or (len(digits) == 11 and digits[0] in "01")
                or (len(digits) == 12 and digits.startswith("91"))
            )
            if not plausible:
                continue

        line_start = text.rfind("\n", 0, m.start()) + 1
        if PHONE_ID_LABEL_RE.search(text[line_start:m.start()]):
            continue
        return candidate
    return None


def extract_education(text: str, mode: str = "highest") -> str:
    """Degree level mentioned in `text` (in the Education section if the resume
    has one, otherwise anywhere in the text).

    mode="highest" - a candidate's level (resumes).
    mode="lowest"  - the minimum a job description asks for, so
                     "Bachelor's required, Master's preferred" -> Bachelor's.
    """
    scopes = []
    edu_text = _section_text(text, "education")
    if edu_text:
        scopes.append(edu_text)
    scopes.append(text)
    levels = EDUCATION_LEVELS if mode == "highest" else list(reversed(EDUCATION_LEVELS))
    for scope in scopes:
        for pattern, label in levels:
            if pattern.search(scope):
                return label
    return "Not detected"


def extract_social_links(text: str) -> dict:
    """Extract the main professional links from resume text."""
    links = []
    for match in URL_RE.finditer(text):
        link = match.group(0).rstrip(".,;:)]}")
        if not link.lower().startswith(("http://", "https://")):
            link = "https://" + link
        if link not in links:
            links.append(link)

    social = {"linkedin_url": "Not detected", "github_url": "Not detected", "portfolio_url": "Not detected"}
    remaining = []
    for link in links:
        lower = link.lower()
        if "linkedin.com" in lower and social["linkedin_url"] == "Not detected":
            social["linkedin_url"] = link
        elif "github.com" in lower and social["github_url"] == "Not detected":
            social["github_url"] = link
        else:
            remaining.append(link)
    if remaining:
        social["portfolio_url"] = remaining[0]
    return social


def extract_social_link_labels(text: str) -> dict:
    """Find the resume label associated with each extracted social link."""
    labels = {
        "linkedin_url": "LinkedIn",
        "github_url": "GitHub",
        "portfolio_url": "Portfolio",
    }
    label_re = re.compile(
        r"(linkedin|github|leetcode|portfolio|personal\s+website|website|web|site)\s*(?:and\s+links)?\s*:?\s*$",
        re.IGNORECASE,
    )
    for match in URL_RE.finditer(text):
        link = match.group(0).rstrip(".,;:)]}")
        normalized = link if link.lower().startswith(("http://", "https://")) else "https://" + link
        lower = normalized.lower()
        if "linkedin.com" in lower:
            field = "linkedin_url"
        elif "github.com" in lower:
            field = "github_url"
        else:
            field = "portfolio_url"
        line_start = text.rfind("\n", 0, match.start()) + 1
        prefix = text[line_start:match.start()]
        label_match = label_re.search(prefix)
        if "leetcode.com" in lower:
            labels["portfolio_url"] = "LeetCode"
        elif label_match:
            label = label_match.group(1).strip().lower()
            labels[field] = {
                "linkedin": "LinkedIn",
                "github": "GitHub",
                "leetcode": "LeetCode",
                "portfolio": "Portfolio",
                "personal website": "Personal Website",
                "website": "Website",
                "web": "Web",
                "site": "Site",
            }.get(label, label.title())
    return labels


_JS_FRAMEWORK_PREFIXES = (
    "node", "express", "react", "vue", "next", "nuxt",
    "three", "chart", "d3", "angular", "ember", "backbone",
    "electron", "meteor", "knockout", "alpine", "nest", "remix", "gatsby", "svelte",
)


def _is_standalone_js(text_low: str) -> bool:
    """True only if 'js' appears as a standalone skill token (e.g. 'Languages: JS, Python'),
    not as a file extension/suffix ('.js', 'node.js', 'express.js') and not as part of
    a framework name like 'node js', 'express js', 'react js'."""
    for m in re.finditer(r"(?<![a-z0-9./\-])js(?![a-z0-9])", text_low):
        prefix = text_low[:m.start()].rstrip()
        if any(
            prefix.endswith(f) and (len(prefix) == len(f) or not prefix[-len(f)-1].isalnum())
            for f in _JS_FRAMEWORK_PREFIXES
        ):
            continue
        return True
    return False


def extract_skills(text: str, extra_skills: list = None) -> list:
    """extra_skills lets callers extend the built-in taxonomy at runtime
    (e.g. org-specific tools like LangGraph or vLLM added via Settings)
    without editing skills_taxonomy.py."""
    low = " " + re.sub(r"[^a-z0-9.+#/\s]", " ", text.lower()) + " "
    found = set()

    for skill in MASTER_SKILLS:
        s_low = skill.lower()
        # Skills starting with a dot (e.g. '.NET') allow non-alphanumeric lookbehind;
        # other skills must not be preceded by a dot to prevent matching file extensions/suffixes.
        lb = r"(?<![a-z0-9])" if s_low.startswith(".") else r"(?<![a-z0-9.])"

        # Skills ending with '+' or '#' (e.g. 'C++', 'C#') allow non-alphanumeric lookahead;
        # 'C' alone must not be followed by '+' or '#' so 'C++' or 'C#' doesn't match 'C'.
        if s_low.endswith("+") or s_low.endswith("#"):
            la = r"(?![a-z0-9])"
        elif s_low == "c":
            la = r"(?![a-z0-9+#])"
        else:
            la = r"(?![a-z0-9])"

        pattern = lb + re.escape(s_low) + la
        if re.search(pattern, low):
            found.add(skill)

    for alias, canonical in SYNONYMS.items():
        if alias == "js":
            # 'js' is only JavaScript if standalone; never when part of '.js' or framework names
            if _is_standalone_js(low):
                found.add("JavaScript")
            continue
        lb = r"(?<![a-z0-9])" if alias.startswith(".") else r"(?<![a-z0-9.])"
        la = r"(?![a-z0-9])"
        pattern = lb + re.escape(alias) + la
        if re.search(pattern, low):
            found.add(canonical)

    for skill in (extra_skills or []):
        skill = skill.strip()
        if not skill:
            continue
        s_low = skill.lower()
        lb = r"(?<![a-z0-9])" if s_low.startswith(".") else r"(?<![a-z0-9.])"
        la = r"(?![a-z0-9])"
        pattern = lb + re.escape(s_low) + la
        if re.search(pattern, low):
            found.add(skill)

    return sorted(found)


def parse_document(filename: str, file_bytes: bytes, extra_skills: list = None) -> ParsedDocument:
    text = extract_text(filename, file_bytes)
    fallback_name = re.sub(r"\.[a-zA-Z0-9]+$", "", filename).replace("_", " ").replace("-", " ").title()

    email_match = EMAIL_RE.search(text)
    phone = extract_phone(text)
    social_links = extract_social_links(text)

    return ParsedDocument(
        raw_text=text,
        name=guess_name(text, fallback_name),
        email=email_match.group(0) if email_match else "Not detected",
        phone=phone if phone else "Not detected",
        **social_links,
        education=extract_education(text),
        years_experience=extract_years_experience(text),
        skills=extract_skills(text, extra_skills=extra_skills),
    )


def parse_job_description(jd_text: str, min_years_override=None, extra_skills: list = None) -> dict:
    skills = extract_skills(jd_text, extra_skills=extra_skills)
    years_match = YEARS_EXP_RE.search(jd_text)
    min_years = float(years_match.group(1)) if years_match else (min_years_override or 0.0)
    education = extract_education(jd_text, mode="lowest")
    return {
        "raw_text": jd_text,
        "required_skills": skills,
        "min_years": min_years,
        "required_education": education,
    }


# ---------------------------------------------------------------------------
# Compatibility Helpers
# ---------------------------------------------------------------------------
def format_experience(years, months=None, compact=False) -> str:
    """Format experience displaying months for short periods or years."""
    if months is not None and months > 0:
        if months < 12:
            return f"{months} mo{'s' if months != 1 else ''}" if compact else f"{months} month{'s' if months != 1 else ''}"
        y = months // 12
        rem_m = months % 12
        if rem_m == 0:
            return f"{y} yr{'s' if y != 1 else ''}" if compact else f"{y} year{'s' if y != 1 else ''}"
        return f"{y} yr{'s' if y != 1 else ''} {rem_m} mo{'s' if rem_m != 1 else ''}" if compact else f"{y} year{'s' if y != 1 else ''} {rem_m} month{'s' if rem_m != 1 else ''}"
    try:
        y_val = float(years or 0)
    except (TypeError, ValueError):
        y_val = 0.0
    y_str = f"{y_val:.1f}"
    if y_str.endswith(".0"):
        y_str = y_str[:-2]
    if compact:
        suffix = "yr" if y_str == "1" else "yrs"
        return f"{y_str} {suffix}"
    suffix = "year" if y_str == "1" else "years"
    return f"{y_str} {suffix}"


def extract_experience_months(text: str, now=None) -> int:
    try:
        now = now or _dt.date.today()
        return _merged_months([iv for iv, _, _ in _work_periods(text, now)])
    except Exception:
        return 0


def clean_doubled_text(text: str) -> str:
    if not text:
        return ""
    def _fix_word(match):
        w = match.group(0)
        if len(w) >= 6 and len(w) % 2 == 0:
            if all(w[i].lower() == w[i + 1].lower() for i in range(0, len(w), 2)):
                return "".join(w[i] for i in range(0, len(w), 2))
        return w
    return re.sub(r"\b[A-Za-z]{6,}\b", _fix_word, text)


def extract_email(text: str) -> str:
    m = EMAIL_RE.search(text or "")
    return m.group(0) if m else "Not detected"


def compute_parse_confidence(doc, text: str) -> tuple:
    return 1.0, []
